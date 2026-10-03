#!/usr/bin/env python3
"""Krabi (KBV): cost depends on hotel x season x package x day-3, so it is modelled with
page selectors (destination "options") instead of one TO per combination.

R&D (PT KRABI R&D - reformatupload.xlsx):
  CR rows 4-25: Star | Hotel | Season | Key | Budget THB | Std·Emerald | Std·PhiPhi | HNY·Emerald | HNY·PhiPhi
  Cost/pax  = CR!B2 (FX) x rate[hotel|season][column]      (flat for every pax)
  Catalog   = per package and pax (Costing col D) + upgrade by hotel tier (H3):
              4★ +250 (Honeymoon +298), 4★+ +550 (Honeymoon +598)
  Selling   = catalog - discount tier 2 (F2); CWB/CNB cost = adult x B3 / D3.

Every hotel x season x package x day-3 combination is recalculated by LibreOffice and
written to <work>/truth/KBV.json; build verifies the page model against all of them.

    python3 tools/extract_krabi.py --work <dir>      # writes <dir>/dest/KBV.json
"""
import argparse, json, os, re, sys
import openpyxl

sys.path.insert(0, os.path.dirname(__file__))
from extract_rd import RD_ROOT, pick_file, sheet, recalc, costing_rows, slug, num, PO, COUNTRY
from build_data import addons

COLS = [("bud", "Budget"), ("stdA", "Standard · Emerald"), ("stdC", "Standard · Phi Phi"),
        ("hnyA", "Honeymoon · Emerald"), ("hnyC", "Honeymoon · Phi Phi")]
# variant id -> (package, day-3 selector value in Costing!H2, CR column)
VARIANTS = [("budget", "Budget", "Emerald", "bud", "Budget (Pkg B)"),
            ("std-emerald", "Standard", "Emerald", "stdA", "Day 3: Emerald Pool + Tiger Cave (Pkg A)"),
            ("std-phiphi", "Standard", "Phi Phi", "stdC", "Day 3: Phi Phi + Maya Bay speedboat (Pkg C)"),
            ("hny-emerald", "Honeymoon", "Emerald", "hnyA", "Day 3: Emerald Pool + Tiger Cave (Pkg A)"),
            ("hny-phiphi", "Honeymoon", "Phi Phi", "hnyC", "Day 3: Phi Phi + Maya Bay speedboat (Pkg C)")]
UPGRADE = {"4★": {"budget": 250, "standard": 250, "honeymoon": 298},
           "4★+": {"budget": 550, "standard": 550, "honeymoon": 598}}


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--work", required=True); a = ap.parse_args()
    work = os.path.abspath(a.work); profile = os.path.join(work, "lo-profile")
    src = pick_file("KRABI")
    wb = openpyxl.load_workbook(src)
    cr = sheet(wb, "CR")
    fx = cr["B2"].value
    hotels, rates = [], []
    for r in range(4, 26):
        star, hotel, season = cr.cell(r, 1).value, cr.cell(r, 2).value, cr.cell(r, 3).value
        if not hotel:
            continue
        hid = slug(hotel).lower()
        if not any(h["id"] == hid for h in hotels):
            up = UPGRADE.get(star, {})
            hotels.append({"id": hid, "label": "%s (%s)" % (hotel, star), "name": hotel, "star": star,
                           **({"upgrade": up} if up else {})})
        for k, (cid, clabel) in enumerate(COLS):
            rates.append({"id": "%s|%s|%s" % (hid, season, cid), "label": "%s · %s · %s" % (hotel, season, clabel),
                          "value": cr.cell(r, 5 + k).value, "fx": "THB", "unit": "THB/pax", "group": hotel})
    seasons = sorted({r["id"].split("|")[1] for r in rates}, key=lambda s: ["Normal", "High"].index(s) if s in ("Normal", "High") else 9)

    # recalc every combination with the workbook's own formulas
    combos, copies = [], []
    base = os.path.join(work, "krabi"); os.makedirs(base, exist_ok=True)
    for h in hotels:
        for season in seasons:
            for vid, pkg, day3, col, _ in VARIANTS:
                w = openpyxl.load_workbook(src); cs = sheet(w, "Costing")
                cs["C2"].value, cs["H2"].value, cs["J2"].value, cs["J3"].value = pkg, day3, h["name"], season
                p = os.path.join(base, "KBV__%s__%s__%s.xlsx" % (h["id"], season, vid)); w.save(p)
                combos.append((h["id"], season, vid, os.path.basename(p))); copies.append(p)
    recalc(copies, os.path.join(base, "out"), profile)
    truth, pricing = [], {}
    for hid, season, vid, fn in combos:
        cs = sheet(openpyxl.load_workbook(os.path.join(base, "out", fn), data_only=True), "Costing")
        rows = costing_rows(cs)
        truth.append({"options": {"hotel": hid, "season": season}, "variant": vid,
                      "rows": {k: {str(p): list(v) for p, v in rows.get(k, {}).items()} for k in ("adult", "cwb", "cnb")},
                      "upgrade": cs["H3"].value, "disc": cs["F2"].value, "r3": [c.value for c in cs[3]][:4]})
        pk = {"budget": "budget", "std": "standard", "hny": "honeymoon"}[vid.split("-")[0]]
        if pk not in pricing and not (cs["H3"].value or 0):  # base catalog from a hotel without upgrade
            pricing[pk] = {k: {str(p): v[1] for p, v in rows.get(k, {}).items() if num(v[1])} for k in ("adult", "cwb", "cnb")}
            pricing[pk]["rules"] = {"discountTier2": cs["F2"].value, "cwb": cs["B3"].value, "cnb": cs["D3"].value}
    variants = [{"id": vid, "label": "%s · %s" % (pkg, lbl) if vid != "budget" else lbl, "supplier": "KTT Travels",
                 "paxMin": 2, "paxMax": 30, "hotel": "Selected hotel", "notes": "",
                 "components": [{"key": "hotel-package", "label": "KTT package (hotel + tours)", "per": "pax",
                                 "expr": "R[O.hotel+'|'+O.season+'|%s']" % col}]}
                for vid, pkg, day3, col, lbl in VARIANTS]
    packages = []
    for pid, label, assign in (("budget", "Krabi Budget 4D3N", ["budget"]),
                               ("standard", "Krabi Standard 4D3N", ["std-emerald"]),
                               ("honeymoon", "Krabi Honeymoon 4D3N", ["hny-emerald"])):
        pr = pricing[pid]
        packages.append({"id": pid, "label": label, "catalog": None,
                         "assign": [{"from": 2, "to": 30, "variant": assign[0]}],
                         "pricing": {"adult": pr["adult"], "cwb": pr["cwb"], "cnb": pr["cnb"], "infant": 0,
                                     "source": {k: {p: "rd" for p in pr[k]} for k in ("adult", "cwb", "cnb")}},
                         "rules": {"discountTier2": pr["rules"]["discountTier2"], "tierUpgrade": 0,
                                   "cwbCost": {"type": "pct", "value": pr["rules"]["cwb"]},
                                   "cnbCost": {"type": "pct", "value": pr["rules"]["cnb"]},
                                   "infantCost": {"type": "flat", "value": 0}}})
    dest = {"code": "KBV", "name": "Krabi", "country": COUNTRY["KBV"], "duration": "4D3N", "nights": 3,
            "po": PO["KBV"], "source": os.path.relpath(src, RD_ROOT),
            "note": "Cost = selected hotel's KTT rate (THB) × FX, flat per pax. Hotel tier adds the catalog upgrade. Standard and Honeymoon sell at the same price for either Day 3; the Tour operator box picks Day 3.",
            "options": [{"id": "hotel", "label": "Hotel", "choices": hotels},
                        {"id": "season", "label": "Season", "choices": [{"id": s, "label": s} for s in seasons]}],
            "fx": [{"id": "MYR", "label": "MYR", "value": 1, "locked": True}, {"id": "THB", "label": "THB → MYR", "value": fx}],
            "rates": rates, "tables": [], "variants": variants, "packages": packages, "addons": addons(wb, "KBV")}
    os.makedirs(os.path.join(work, "dest"), exist_ok=True); os.makedirs(os.path.join(work, "truth"), exist_ok=True)
    json.dump(dest, open(os.path.join(work, "dest", "KBV.json"), "w"), indent=1, ensure_ascii=False)
    json.dump({"code": "KBV", "combos": truth}, open(os.path.join(work, "truth", "KBV.json"), "w"), default=str)
    # verify in Python the way the page computes it
    rate = {r["id"]: r["value"] * fx for r in rates}
    hmap = {h["id"]: h for h in hotels}
    bad = n = 0
    for t in truth:
        vid = t["variant"]; pk = {"budget": "budget", "std": "standard", "hny": "honeymoon"}[vid.split("-")[0]]
        P = next(p for p in packages if p["id"] == pk)
        col = next(c for v, _, _, c, _ in VARIANTS if v == vid)
        a = rate["%s|%s|%s" % (t["options"]["hotel"], t["options"]["season"], col)]
        up = hmap[t["options"]["hotel"]].get("upgrade", {}).get(pk, 0)
        if abs(up - (t["upgrade"] or 0)) > 0.01:
            bad += 1; print("upgrade", t["options"], vid, up, t["upgrade"])
        for k, mult in (("adult", 1), ("cwb", P["rules"]["cwbCost"]["value"]), ("cnb", P["rules"]["cnbCost"]["value"])):
            for p, (c, cat, sell) in t["rows"][k].items():
                n += 1
                mine_sell = P["pricing"][k][p] + up - P["rules"]["discountTier2"]
                if abs(a * mult - c) > 0.6 or abs(mine_sell - sell) > 0.6:
                    bad += 1
                    if bad < 5: print("diff", t["options"], vid, k, p, a * mult, c, mine_sell, sell)
    print("KBV: %d hotels x %d seasons x %d variants; verified %d rows, %d mismatches" % (len(hotels), len(seasons), len(VARIANTS), n, bad))


if __name__ == "__main__":
    main()
