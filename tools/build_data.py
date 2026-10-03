#!/usr/bin/env python3
"""Build data/data.json for the PT costing calculator from the R&D workbooks
and the Catalog PT site JSON.

Selling price  = Catalog PT site (catalogs/<slug>.json) for every pax band the
                 catalog prints; pax beyond the catalog fall back to the R&D
                 Costing tab and are tagged source "rd".
Cost           = R&D CR tab, rebuilt as editable components (rates x qty).

Run once to seed the repo. After editors start saving from the page, the
data.json in the GitHub repo is the source of truth -- do not overwrite it
with a fresh build unless you mean to discard their edits.

    python3 tools/build_data.py            # writes data/data.json
    python3 tools/build_data.py --check    # report R&D vs catalog differences only
"""
import json, re, sys, datetime, os
import openpyxl

DRIVE = os.environ.get("PT_DRIVE", os.path.expanduser(
    "~/Library/CloudStorage/GoogleDrive-product@arbatravel.com/My Drive"))
CATALOG = os.path.join(DRIVE, "Catalog PT (new)", "catalogs")
OUT = os.path.join(os.path.dirname(__file__), "..", "data", "data.json")

warnings = []


def wb(name, data_only=False):
    return openpyxl.load_workbook(os.path.join(DRIVE, name), data_only=data_only)


# ---------------------------------------------------------------- catalog
def expand_pax(label):
    nums = [int(n) for n in re.findall(r"\d+", label)]
    return list(range(nums[0], nums[-1] + 1)) if len(nums) > 1 else nums


def money(s):
    return float(re.sub(r"[^\d.]", "", s))


def catalog_prices(slug):
    d = json.load(open(os.path.join(CATALOG, slug + ".json")))
    out = {"adult": {}, "cwb": {}, "cnb": {}}
    for row in d["prices"]["rows"]:
        for p in expand_pax(row["pax"]):
            for k, amt in zip(("adult", "cwb", "cnb"), row["amounts"]):
                out[k][p] = money(amt)
    infant = money(d["prices"].get("infant", "0") or "0")
    meta = {"slug": slug, "title": d["title"], "version": d.get("version"),
            "updated": d.get("updated"), "valid_until": d.get("valid_until"),
            "url": "https://prod-at22.github.io/catalog-pt-public/%s.html" % slug}
    return out, infant, meta


# ---------------------------------------------------------------- R&D catalog (IF chains)
def parse_if(formula, names):
    """=IF($C$2="A",1,IF($C$2="B",2,3))+$H$3 -> {A:1,B:2,<other>:3}"""
    if isinstance(formula, (int, float)):
        return {n: float(formula) for n in names}
    s = str(formula).replace("$", "")
    res = {}
    for n, v in re.findall(r'C2="([^"]+)",("—"|[\d.]+)', s):
        res[n] = None if v == '"—"' else float(v)
    m = re.search(r'=(\d+(?:\.\d+)?)\+H3$', s)  # plain "=4397+$H$3"
    tail = re.findall(r',("—"|[\d.]+)\)', s)
    other = m.group(1) if m else (tail[-1] if tail else None)
    for n in names:
        if n not in res:
            res[n] = None if other in (None, '"—"') else float(other)
    return res


def rd_catalog(book, names, starts=(6, 37, 68)):
    ws = book["Costing"]
    out = {k: {n: {} for n in names} for k in ("adult", "cwb", "cnb")}
    for sec, start in zip(("adult", "cwb", "cnb"), starts):
        for i in range(29):
            pax = ws.cell(start + i, 1).value
            if not isinstance(pax, (int, float)):
                continue
            for n, v in parse_if(ws.cell(start + i, 4).value, names).items():
                if v is not None:
                    out[sec][n][int(pax)] = v
    return out


def merge_pricing(cat, infant, rd_by_cat, label):
    """Catalog first, R&D fills pax the catalog does not print."""
    pricing = {"infant": infant, "source": {}}
    for k in ("adult", "cwb", "cnb"):
        merged = {}
        for p, v in rd_by_cat[k].items():
            merged[p] = v
            pricing["source"].setdefault(k, {})[str(p)] = "rd"
        for p, v in cat.get(k, {}).items():
            if p in rd_by_cat[k] and abs(rd_by_cat[k][p] - v) > 0.5:
                warnings.append("%s %s pax %d: catalog RM%.0f vs R&D RM%.0f (catalog used)"
                                % (label, k.upper(), p, v, rd_by_cat[k][p]))
            merged[p] = v
            pricing["source"].setdefault(k, {})[str(p)] = "catalog"
        pricing[k] = {str(p): merged[p] for p in sorted(merged)}
    return pricing


# ---------------------------------------------------------------- SEOUL
def seoul():
    b = wb("PT_SEL_RD_reformatted.xlsx")
    cr = b["CR"]
    keta, profit_usd, usd = cr["D3"].value, cr["F3"].value, cr["H3"].value
    tables, variants, ranges, hotels, notes = [], [], {}, {}, {}
    for row in cr.iter_rows(min_row=6, values_only=True):
        if isinstance(row[0], (int, float)) and row[1]:
            to = row[1]
            ranges.setdefault(to, []).append(int(row[0]))
            hotels[to] = row[7]
            notes.setdefault(to, row[8])
            t = next((t for t in tables if t["id"] == to), None)
            if not t:
                t = {"id": to, "label": "Ground cost " + to, "unit": "MYR/pax",
                     "fx": "MYR", "group": "ATK package rate (Ground)", "values": {}}
                tables.append(t)
            t["values"][str(int(row[0]))] = row[2]
    names = list(ranges)
    labels = {"ATK-PT-BSC": "ATK · PT Basic", "ATK-PT-STD": "ATK · PT Standard",
              "ATK-ST": "ATK · Self Tour (estimated)"}
    for to in names:
        variants.append({
            "id": to, "label": labels.get(to, to), "supplier": "ATK (ARBA Travel Korea)",
            "paxMin": min(ranges[to]), "paxMax": max(ranges[to]),
            "hotel": hotels[to], "notes": notes[to],
            "components": [
                {"key": "ground", "label": "Ground (ATK package rate)", "per": "pax",
                 "expr": "T['%s']" % to},
                {"key": "tipping", "label": "Tipping", "per": "pax", "expr": "R.tipping"},
                {"key": "keta", "label": "K-ETA", "per": "pax", "expr": "R.keta"},
                {"key": "profit", "label": "ATK profit", "per": "pax", "expr": "R.atkProfit"},
            ]})
    rd = rd_catalog(b, names)
    rd_for = lambda to: {k: rd[k][to] for k in rd}
    rules = {"cwbCost": {"type": "pct", "value": b["Costing"]["B3"].value},
             "cnbCost": {"type": "pct", "value": b["Costing"]["D3"].value},
             "infantCost": {"type": "flat", "value": b["Costing"]["F3"].value},
             "discountTier2": b["Costing"]["F2"].value, "tierUpgrade": b["Costing"]["H3"].value}
    pk = []
    for pid, label, slug, to in (("basic", "Seoul Basic 5D4N", "seoul-basic", "ATK-PT-BSC"),
                                 ("standard", "Seoul Standard 5D4N", "seoul-standard", "ATK-PT-STD"),
                                 ("selftour", "Seoul Self Tour 5D4N", None, "ATK-ST")):
        if slug:
            cat, inf, meta = catalog_prices(slug)
        else:
            cat, inf, meta = {}, 200.0, None
        pk.append({"id": pid, "label": label, "catalog": meta,
                   "assign": [{"from": min(ranges[to]), "to": max(ranges[to]), "variant": to}],
                   "pricing": merge_pricing(cat, inf, rd_for(to), label),
                   "rules": json.loads(json.dumps(rules))})
    return {
        "code": "SEL", "name": "Seoul", "country": "Korea", "duration": "5D4N",
        "nights": 4, "source": "PT_SEL_RD_reformatted.xlsx",
        "note": "CR note: ATK package rate usually already includes ATK profit. If so, set ATK profit (USD) to 0 to avoid double-counting.",
        "fx": [{"id": "MYR", "label": "MYR", "value": 1, "locked": True},
               {"id": "USD", "label": "USD → MYR", "value": usd}],
        "rates": [
            {"id": "tipping", "label": "Tipping (bundled in ground)", "value": 0, "fx": "MYR", "unit": "MYR/pax", "group": "Per-pax add-ons"},
            {"id": "keta", "label": "K-ETA", "value": keta, "fx": "MYR", "unit": "MYR/pax", "group": "Per-pax add-ons"},
            {"id": "atkProfit", "label": "ATK profit", "value": profit_usd, "fx": "USD", "unit": "USD/pax", "group": "Per-pax add-ons"},
        ],
        "tables": tables, "variants": variants, "packages": pk}


# ---------------------------------------------------------------- TOKYO
def tokyo():
    b = wb("PT_HND_RD_reformatted.xlsx")
    rc, cr = b["Raw Costing"], b["CR"]
    c = lambda cell: rc[cell].value
    R = lambda id, cell, fx, unit, group: {"id": id, "label": rc["A" + cell[1:]].value + " (" + rc["B" + cell[1:]].value + ")",
                                           "value": c(cell), "fx": fx, "unit": unit, "group": group}
    rates = [
        R("hnd7", "C6", "WIF", "JPY/way", "1. Airport transfer"),
        R("hnd10", "C7", "WIF", "JPY/way", "1. Airport transfer"),
        R("hnd14", "C8", "WIF", "JPY/way", "1. Airport transfer"),
        R("hndBig", "C9", "WIF", "JPY/way", "1. Airport transfer"),
        R("hndQ", "C10", "QAYYUM", "JPY/way", "1. Airport transfer"),
        R("qCity", "C14", "QAYYUM", "JPY/trip", "2. Tour transport"),
        R("qFuji", "C15", "QAYYUM", "JPY/trip", "2. Tour transport"),
        R("wCity89", "C16", "WIF", "JPY/trip", "2. Tour transport"),
        R("wCity1012", "C17", "WIF", "JPY/trip", "2. Tour transport"),
        R("wFuji89", "C18", "WIF", "JPY/trip", "2. Tour transport"),
        R("wFuji1012", "C19", "WIF", "JPY/trip", "2. Tour transport"),
        R("ptPass", "C23", "WIF", "JPY/pax/day", "3. Public transport"),
        R("hwBus", "C24", "WIF", "JPY/pax/way", "3. Public transport"),
        R("hotel3", "C28", "MYR", "MYR/pax/night", "4. Accommodation"),
        R("apt", "C29", "MYR", "MYR/pax/night", "4. Accommodation"),
        R("stAccom", "C30", "WIF", "JPY/pax/night", "4. Accommodation"),
        R("iyashi", "C34", "WIF", "JPY/pax", "5. Entrance"),
        R("kachi", "C35", "WIF", "JPY/pax", "5. Entrance"),
        R("gFull", "C39", "WIF", "JPY/day", "6. Guide (WIF)"),
        R("gHalf", "C40", "WIF", "JPY/day", "6. Guide (WIF)"),
        R("gAcc", "C41", "WIF", "JPY/day", "6. Guide (WIF)"),
        R("gMeal", "C42", "WIF", "JPY/day", "6. Guide (WIF)"),
        R("gTrn", "C43", "WIF", "JPY/day", "6. Guide (WIF)"),
        R("wifCharge", "C48", "WIF", "JPY/pax", "7. Service charge"),
    ]
    guide = "(R.gHalf+R.gAcc+R.gMeal+R.gTrn)+(R.gFull+R.gAcc+R.gMeal+R.gTrn)*4"
    ranges, hotels, notes = {}, {}, {}
    for row in cr.iter_rows(min_row=6, values_only=True):
        if isinstance(row[0], (int, float)) and row[1]:
            ranges.setdefault(row[1], []).append(int(row[0]))
            hotels[row[1]] = row[13]
            notes.setdefault(row[1], row[14])
    comp = lambda key, label, expr: {"key": key, "label": label, "per": "group", "expr": expr}
    variants = [
        {"id": "WIF-BSC", "label": "WIF · Basic (public transport + guide)", "supplier": "WIF",
         "components": [
             comp("airport", "Airport transfer ×2", "2*band(pax,[4,R.hnd7],[8,R.hnd10],[999,R.hndBig])"),
             comp("transport", "Public transport pass ×N days", "R.ptPass*N*pax"),
             comp("bus", "Highway bus Tokyo–Fuji ×2", "R.hwBus*2*pax"),
             comp("entrance", "Entrance (Kachi Kachi Ropeway)", "R.kachi*pax"),
             comp("accomm", "Hotel 3★ ×N nights", "R.hotel3*N*pax"),
             comp("guide", "Tour guide (½ day + 4 full days)", guide),
             comp("charge", "WIF service charge", "R.wifCharge*pax")]},
        {"id": "QAYYUM-STD", "label": "Qayyum · Standard (pax 2–7)", "supplier": "Qayyum",
         "components": [
             comp("airport", "Airport transfer ×2", "2*R.hndQ"),
             comp("transport", "Private transport (Fuji + 3× city)", "R.qFuji+R.qCity*3"),
             comp("entrance", "Entrance (Iyashi No Sato)", "R.iyashi*pax"),
             comp("accomm", "Apartment ×N nights", "R.apt*N*pax")]},
        {"id": "WIF-STD", "label": "WIF · Standard (pax 8+)", "supplier": "WIF",
         "components": [
             comp("airport", "Airport transfer ×2", "2*band(pax,[9,R.hnd10],[999,R.hndBig])"),
             comp("transport", "Private transport (Fuji + 3× city)", "band(pax,[9,R.wFuji89+R.wCity89*3],[999,R.wFuji1012+R.wCity1012*3])"),
             comp("entrance", "Entrance (Iyashi No Sato)", "R.iyashi*pax"),
             comp("accomm", "Apartment ×N nights", "R.apt*N*pax"),
             comp("guide", "Tour guide (½ day + 4 full days)", guide),
             comp("charge", "WIF service charge", "R.wifCharge*pax")]},
        {"id": "WIF-ST", "label": "WIF · Self Tour", "supplier": "WIF",
         "components": [
             comp("airport", "Airport transfer ×2", "2*band(pax,[3,R.hnd7],[7,R.hnd10],[999,R.hnd14])"),
             comp("accomm", "Customer accommodation ×N nights", "R.stAccom*N*pax"),
             comp("charge", "WIF service charge", "R.wifCharge*pax")]},
    ]
    for v in variants:
        v.update({"paxMin": min(ranges[v["id"]]), "paxMax": max(ranges[v["id"]]),
                  "hotel": hotels[v["id"]], "notes": notes[v["id"]]})
    names = list(ranges)
    rd = rd_catalog(b, names)
    cs = b["Costing"]
    rules = {"cwbCost": {"type": "pct", "value": cs["B3"].value},
             "cnbCost": {"type": "minus", "value": cs["D3"].value},
             "infantCost": {"type": "flat", "value": 0},
             "discountTier2": cs["F2"].value, "tierUpgrade": cs["H3"].value}
    std_rd = {k: {**rd[k]["QAYYUM-STD"], **rd[k]["WIF-STD"]} for k in rd}
    pk = []
    cat, inf, meta = catalog_prices("tokyo-basic")
    pk.append({"id": "basic", "label": "Tokyo Basic 5D4N", "catalog": meta,
               "assign": [{"from": 2, "to": 40, "variant": "WIF-BSC"}],
               "pricing": merge_pricing(cat, inf, {k: rd[k]["WIF-BSC"] for k in rd}, "Tokyo Basic"),
               "rules": json.loads(json.dumps(rules))})
    cat, inf, meta = catalog_prices("tokyo-standard")
    pk.append({"id": "standard", "label": "Tokyo Standard 5D4N", "catalog": meta,
               "assign": [{"from": 2, "to": 7, "variant": "QAYYUM-STD"},
                          {"from": 8, "to": 40, "variant": "WIF-STD"}],
               "pricing": merge_pricing(cat, inf, std_rd, "Tokyo Standard"),
               "rules": json.loads(json.dumps(rules))})
    pk.append({"id": "selftour", "label": "Tokyo Self Tour 5D4N", "catalog": None,
               "assign": [{"from": 2, "to": 40, "variant": "WIF-ST"}],
               "pricing": merge_pricing({}, 200.0, {k: rd[k]["WIF-ST"] for k in rd}, "Tokyo Self Tour"),
               "rules": json.loads(json.dumps(rules))})
    return {
        "code": "HND", "name": "Tokyo", "country": "Jepun", "duration": "5D4N",
        "nights": 4, "source": "PT_HND_RD_reformatted.xlsx",
        "note": "Standard: Qayyum for 2–7 pax, WIF for 8+. JPY rates convert at their supplier FX; accommodation is in MYR.",
        "fx": [{"id": "MYR", "label": "MYR", "value": 1, "locked": True},
               {"id": "WIF", "label": "WIF JPY → MYR", "value": cr["B3"].value},
               {"id": "QAYYUM", "label": "Qayyum JPY → MYR", "value": cr["B4"].value}],
        "rates": rates, "tables": [], "variants": variants, "packages": pk}


def main():
    dests = [seoul(), tokyo()]
    for w in warnings:
        print("DIFF", w)
    if "--check" in sys.argv:
        return
    now = datetime.datetime.now().astimezone().isoformat(timespec="seconds")
    data = {"schema": 1, "version": 1, "updatedAt": now, "updatedBy": "system",
            "settings": {"marginWarnPct": 0.15, "marginDangerPct": 0.05},
            "destinations": dests}
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    json.dump(data, open(OUT, "w"), indent=1, ensure_ascii=False)
    hist = os.path.join(os.path.dirname(OUT), "history.json")
    json.dump({"entries": [{"v": 1, "at": now, "by": "system",
                            "note": "Initial import: costs from R&D CR tabs, selling prices from Catalog PT site",
                            "changes": []}]}, open(hist, "w"), indent=1, ensure_ascii=False)
    users = os.path.join(os.path.dirname(OUT), "users.json")
    if not os.path.exists(users):
        json.dump({"repo": "", "users": []}, open(users, "w"), indent=1)
    print("wrote", os.path.relpath(OUT), "with", len(dests), "destinations;", len(warnings), "catalog/R&D differences")


if __name__ == "__main__":
    main()
