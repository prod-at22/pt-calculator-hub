#!/usr/bin/env python3
"""Osaka (OSK) and Tokyo-Osaka (KIX): turn the per-pax CR tables into rate-by-rate formulas, like Tokyo,
so Show calculation / Rate reference work on them.

Reads the R&D CR tab FORMULAS (each cell = Raw Costing rate × pax / days × FX cell) and writes, per TO and
component, one expression over R.<rate> (a Raw Costing row at the FX the CR multiplies it by). A formula
that changes with pax becomes band(pax, …). Rates the PO changed on the page are kept (OVERRIDE).
Every component at every pax must equal the current calculator value, or nothing is written.

    python3 tools/build_jp_rates.py [--by product] [--dry]
"""
import argparse, datetime, json, os, re, sys
import openpyxl

sys.path.insert(0, os.path.dirname(__file__))
from extract_rd import RD_ROOT, pick_file

ROOT = os.path.join(os.path.dirname(__file__), "..")
DESTS = {"OSK": "OSAKA", "KIX": "TOKYO OSAKA"}
# Raw Costing row -> value the PO set on the page (v24: accommodation RM250/pax/night; R&D still RM300)
OVERRIDE = {("OSK", 22): 250, ("OSK", 23): 250, ("KIX", 30): 250}
RC = re.compile(r"'Raw Costing'!\$C\$(\d+)")
FXC = re.compile(r"\$B\$(\d+)")


def norm(f, row):
    return re.sub(r"(?<![$A-Z])A%d\b" % row, "pax", str(f).lstrip("=")) if isinstance(f, str) else str(f if f is not None else 0)


def build(code, d):
    path = pick_file(os.path.join(RD_ROOT, DESTS[code]))
    wb = openpyxl.load_workbook(path)
    raw, cr = wb["Raw Costing"], wb["CR"]
    fx_by_row = {}
    for f in d["fx"]:
        m = re.search(r"CR!A(\d+)$", f.get("source", ""))
        if m:
            fx_by_row[int(m.group(1))] = f
    heads = [cr.cell(5, c).value for c in range(3, cr.max_column + 1)]
    labels = {c["label"] for v in d["variants"] for c in v["components"]}   # the calculator's component columns
    cols = [i for i, h in enumerate(heads) if h in labels]
    rows = {}   # TO -> [(pax, [formula per component])]
    for r in range(6, cr.max_row + 1):
        p, to = cr.cell(r, 1).value, cr.cell(r, 2).value
        if isinstance(p, (int, float)) and to:
            rows.setdefault(to, []).append((int(p), [norm(cr.cell(r, 3 + i).value, r) for i in cols]))
    rates = {}

    def rate_id(row, fxrow):
        f = fx_by_row.get(fxrow) if fxrow else None
        rid = "r%d%s" % (row, "_" + re.sub(r"\W", "", f["id"]).lower() if f else "")
        if rid not in rates:
            item, to, val, unit = (raw.cell(row, c).value for c in range(1, 5))
            sup = f["id"].split("-", 1)[-1].title().replace("Qayyum/Ucop", "Qayyum/Ucop") if f else ("COMMON" if str(to).upper() == "COMMON" else str(to))
            sup = {"Wif": "WIF"}.get(sup, sup)
            val = OVERRIDE.get((code, row), val)
            note = "PO RM%s on the page, v24; R&D RM%s" % (val, raw.cell(row, 3).value) if (code, row) in OVERRIDE else ""
            rates[rid] = {"id": rid, "label": "%s (%s%s)" % (item.strip(), sup.upper() if sup == "COMMON" else sup, ", " + note if note else ""),
                          "value": val, "fx": f["id"] if f else "MYR", "unit": unit, "group": "R&D Raw Costing row %d" % row}
        return rid

    def expr(f):
        if f in ("0", "None", ""):
            return "0"
        fxs = set(int(x) for x in FXC.findall(f))
        assert len(fxs) <= 1, (code, f)
        fxrow = fxs.pop() if fxs else None
        e = re.sub(r"\*\$B\$\d+", "", f)
        assert "$B$" not in e, (code, f)
        return RC.sub(lambda m: "R." + rate_id(int(m.group(1)), fxrow), e)

    out = {}
    for to, lst in rows.items():
        comps = []
        for i in range(len(cols)):
            seq = []   # [(max pax, expr)] for consecutive pax with the same formula
            for p, fs in lst:
                e = expr(fs[i])
                if seq and seq[-1][1] == e:
                    seq[-1][0] = p
                else:
                    seq.append([p, e])
            seq[-1][0] = 999
            comps.append(seq[0][1] if len(seq) == 1 else "band(pax,%s)" % ",".join("[%d,%s]" % (m, e) for m, e in seq))
        out[to] = comps
    return path, [heads[i] for i in cols], out, rates


def evaluate(d, rates, e, pax):
    fx = {f["id"]: f["value"] for f in d["fx"]}
    R = type("R", (), {r["id"]: r["value"] * (1 if r["fx"] == "MYR" else fx[r["fx"]]) for r in rates.values()})
    band = lambda p, *pairs: next(v for m, v in pairs if p <= m)
    return eval(re.sub(r"band\(pax,", "band(pax,", e).replace("[", "(").replace("]", ")"), {"R": R, "pax": pax, "band": band})


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--by", default="product"); ap.add_argument("--dry", action="store_true")
    a = ap.parse_args()
    dp, hp = os.path.join(ROOT, "data", "data.json"), os.path.join(ROOT, "data", "history.json")
    data, hist = json.load(open(dp)), json.load(open(hp))
    summary, bad = [], []
    for code in DESTS:
        d = next(x for x in data["destinations"] if x["code"] == code)
        path, heads, exprs, rates = build(code, d)
        T = {t["id"]: t for t in d["tables"]}
        fx = {f["id"]: f["value"] for f in d["fx"]}
        for v in d["variants"]:
            new = exprs[v["id"]]
            assert len(new) == len(v["components"]), (code, v["id"], heads)
            for c, e in zip(v["components"], new):
                assert c["label"] == heads[v["components"].index(c)], (c["label"], heads)
                t = T[re.match(r"T\['(.+)'\]", c["expr"]).group(1)]
                for p in range(v["paxMin"], v["paxMax"] + 1):
                    old = t["values"].get(str(p))
                    if old in (None, ""):
                        continue
                    old = old * fx.get(t["fx"], 1) * (p if c["per"] == "pax" else 1)
                    got = evaluate(d, rates, e, p)
                    if abs(got - old) > 0.005:
                        bad.append("%s %s %s pax %d: old %.2f new %.2f (%s)" % (code, v["id"], c["key"], p, old, got, e))
                c["expr"], c["per"] = e, "group"
        used = set(re.findall(r"T\['(.+?)'\]", json.dumps(d["variants"])))
        d["tables"] = [t for t in d["tables"] if t["id"] in used]
        d["rates"] = list(rates.values())
        summary.append("%s: %d TO lines rate by rate from %s (%d rates)" % (d["name"], len(d["variants"]), os.path.basename(path), len(rates)))
        print(summary[-1])
        for v in d["variants"]:
            for c in v["components"]:
                print("   %-14s %-12s %s" % (v["id"], c["key"], c["expr"]))
    if bad:
        print("\n".join(["MISMATCH — nothing written:"] + bad[:40])); sys.exit(1)
    if a.dry:
        print("dry run: all components match"); return
    now = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")
    data["version"] += 1; data["updatedAt"] = now; data["updatedBy"] = a.by
    hist["entries"].append({"v": data["version"], "at": now, "by": a.by, "rebase": True, "changes": [],
                            "note": "Osaka & Tokyo-Osaka rate by rate (Show calculation / Rate reference); same costs",
                            "summary": summary})
    for f, o in ((dp, data), (hp, hist)):
        open(f, "w").write(json.dumps(o, indent=1, ensure_ascii=False) + "\n")
    print("written, data v%d" % data["version"])


if __name__ == "__main__":
    main()
